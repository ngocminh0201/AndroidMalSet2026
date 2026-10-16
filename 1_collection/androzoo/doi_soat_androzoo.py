# -*- coding: utf-8 -*-
"""
Doi soat ba chieu cua tuyen AndroZoo:

    thu muc apk   <->   Ket_qua_AndroZoo.xlsx   <->   androzoo_index.db

Ba noi nay de lech nhau ma khong ai bao:

  a. Dong co trong Excel ma file khong con tren dia. Ket_qua_AndroZoo.xlsx chi
     ghi noi them, khong bao gio doi chieu lai voi o dia, nen file bi xoa di thi
     dong van nam nguyen. Trong DB nhung mau nay van mang trang_thai='da_tai'
     nen tai_androzoo.py se khong bao gio chon lai chung.
  b. File co tren dia ma khong co dong nao trong Excel. Xay ra khi mot dot tai
     bi tat cung (khong phai Ctrl+C) truoc luc ghi Excel: file da tai xong nam
     lai tren dia con dong thi mat.

Script nay chi doi soat va in bang. Muon no sua thi them co:

  --mo-lai         tra trang_thai trong DB ve 'chua' cho cac mau mat file, va
                   ghi danh sach SHA-256 ra sha_can_tai.txt de tai lai bang
                   tai_androzoo.py --danh-sach sha_can_tai.txt
  --bo-sung-dong   lay lai thong tin tu androzoo_index.db va ghi bo sung dong
                   vao Ket_qua_AndroZoo.xlsx cho cac file mo coi

Vi du:

    python doi_soat_androzoo.py
    python doi_soat_androzoo.py --mo-lai --bo-sung-dong
"""

import argparse
import datetime
import os
import re
import sqlite3
import sys

THU_MUC = os.path.dirname(os.path.abspath(__file__))
TEN_EXCEL = "Ket_qua_AndroZoo.xlsx"
TEN_DANH_SACH = "sha_can_tai.txt"
S_MAU = "Mau_AndroZoo"
HEX64 = re.compile(r"^[0-9a-fA-F]{64}$")

COT_DB = ["sha256", "pkg_name", "apk_size", "moc_ngay", "dex_date",
          "vt_detection", "vt_scan_date", "markets", "nhom"]


def doc_tham_so():
    p = argparse.ArgumentParser(
        description="Doi soat thu muc apk voi Ket_qua_AndroZoo.xlsx va "
                    "androzoo_index.db.")
    p.add_argument("--thu-muc", default=os.path.join(THU_MUC, "apk"),
                   help="thu muc chua APK, mac dinh la thu muc apk canh script")
    p.add_argument("--db", default=os.path.join(THU_MUC, "androzoo_index.db"))
    p.add_argument("--excel", default=os.path.join(THU_MUC, TEN_EXCEL))
    p.add_argument("--danh-sach", default=os.path.join(THU_MUC, TEN_DANH_SACH),
                   help="file ghi danh sach SHA-256 can tai lai")
    p.add_argument("--mo-lai", action="store_true",
                   help="tra trang_thai trong DB ve 'chua' cho mau mat file va "
                        "ghi danh sach SHA-256 ra file")
    p.add_argument("--bo-sung-dong", action="store_true",
                   help="ghi bo sung dong vao Excel cho file co tren dia ma "
                        "thieu dong")
    return p.parse_args()


def quet_thu_muc(thu_muc):
    """Tra ve (dict sha -> byte, danh sach file la, danh sach file .tam)."""
    if not os.path.isdir(thu_muc):
        sys.exit("Khong thay thu muc: %s" % thu_muc)
    apk, la, tam = {}, [], []
    for ten in os.listdir(thu_muc):
        dd = os.path.join(thu_muc, ten)
        if not os.path.isfile(dd):
            continue
        goc, duoi = os.path.splitext(ten)
        duoi = duoi.lower()
        if duoi == ".apk" and HEX64.match(goc):
            apk[goc.lower()] = os.path.getsize(dd)
        elif duoi == ".tam":
            tam.append(ten)
        else:
            la.append(ten)
    return apk, la, tam


def doc_excel(dd):
    """Tra ve dict sha -> dict cua sheet Mau_AndroZoo."""
    from openpyxl import load_workbook
    if not os.path.exists(dd):
        sys.exit("Khong thay %s" % dd)
    wb = load_workbook(dd, read_only=True)
    if S_MAU not in wb.sheetnames:
        sys.exit("File %s khong co sheet %s" % (os.path.basename(dd), S_MAU))
    ws = wb[S_MAU]
    hang = ws.iter_rows(values_only=True)
    try:
        tieu_de = next(hang)
    except StopIteration:
        return {}
    ra = {}
    for r in hang:
        if not r or not r[1]:
            continue
        h = str(r[1]).strip().lower()
        if HEX64.match(h):
            ra[h] = dict(zip(tieu_de, r))
    wb.close()
    return ra


def doc_db(con, danh_sach):
    """Tra ve dict sha -> dict theo COT_DB, tra theo tung lo 500 hash."""
    ra = {}
    ds = list(danh_sach)
    for i in range(0, len(ds), 500):
        lo = ds[i:i + 500]
        sql = ("SELECT %s FROM ungvien WHERE sha256 IN (%s)"
               % (",".join(COT_DB), ",".join("?" * len(lo))))
        for hang in con.execute(sql, lo):
            ra[str(hang[0]).lower()] = dict(zip(COT_DB, hang))
    return ra


def gon(n):
    return "{:,}".format(n).replace(",", ".")


def main():
    ts = doc_tham_so()
    if not os.path.exists(ts.db):
        sys.exit("Khong thay danh muc %s" % ts.db)

    apk, la, tam = quet_thu_muc(ts.thu_muc)
    dong = doc_excel(ts.excel)

    thieu_file = sorted(set(dong) - set(apk))
    thieu_dong = sorted(set(apk) - set(dong))

    print("=" * 70)
    print("DOI SOAT TUYEN ANDROZOO")
    print("=" * 70)
    print("  Thu muc apk        : %s" % ts.thu_muc)
    print("  File .apk hop le   : %s" % gon(len(apk)))
    if tam:
        print("  File .tam con sot  : %d (%s)"
              % (len(tam), ", ".join(tam[:3]) + (" ..." if len(tam) > 3 else "")))
    if la:
        print("  File khac trong thu muc: %d (%s)"
              % (len(la), ", ".join(la[:3]) + (" ..." if len(la) > 3 else "")))
    print("  Dong trong %s: %s" % (os.path.basename(ts.excel), gon(len(dong))))
    print()
    print("  Dong co ma file khong con : %s" % gon(len(thieu_file)))
    print("  File co ma thieu dong     : %s" % gon(len(thieu_dong)))
    print("  Khop ca hai phia          : %s" % gon(len(set(dong) & set(apk))))

    if thieu_file:
        theo_lan = {}
        tong_byte = 0
        for h in thieu_file:
            d = dong[h]
            k = str(d.get("Lan_chay") or "khong ro")
            theo_lan[k] = theo_lan.get(k, 0) + 1
            try:
                tong_byte += int(d.get("Kich_thuoc_byte") or 0)
            except (TypeError, ValueError):
                pass
        print()
        print("--- Dong mat file, tach theo lan chay ---")
        for k in sorted(theo_lan):
            print("  %-21s %6d" % (k, theo_lan[k]))
        print("  Tai lai het ton khoang %.1f GB." % (tong_byte / 1024.0 ** 3))

    con = sqlite3.connect(ts.db, timeout=60.0)
    con.execute("PRAGMA busy_timeout=60000")

    if thieu_file:
        trang = {}
        for hang in con.execute(
                "SELECT trang_thai, COUNT(*) FROM ungvien WHERE sha256 IN (%s) "
                "GROUP BY trang_thai" % ",".join("?" * min(len(thieu_file), 900)),
                thieu_file[:900]):
            trang[hang[0]] = hang[1]
        if trang:
            print("  Trang thai trong DB (900 mau dau): %s"
                  % ", ".join("%s=%d" % kv for kv in sorted(trang.items())))

    # --- mo lai trang thai cho mau mat file ---
    if ts.mo_lai and thieu_file:
        ghi_chu = "mo lai do mat file, doi soat %s" % datetime.date.today().strftime("%d/%m/%Y")
        n = 0
        for i in range(0, len(thieu_file), 500):
            lo = thieu_file[i:i + 500]
            cur = con.execute(
                "UPDATE ungvien SET trang_thai='chua', so_lan_thu=0, "
                "thoi_diem_tai=NULL, ghi_chu=? WHERE sha256 IN (%s)"
                % ",".join("?" * len(lo)), [ghi_chu] + lo)
            n += cur.rowcount
        con.commit()
        print()
        print("Da tra %d mau ve trang_thai='chua' trong %s."
              % (n, os.path.basename(ts.db)))
        if n != len(thieu_file):
            print("  Luu y: %d mau khong tim thay trong danh muc, co the danh muc "
                  "da duoc nap lai voi tieu chi khac." % (len(thieu_file) - n))
        with open(ts.danh_sach, "w", encoding="utf-8") as f:
            f.write("\n".join(thieu_file) + "\n")
        print("Da ghi %d dong SHA-256 vao %s"
              % (len(thieu_file), os.path.basename(ts.danh_sach)))
        print("Tai lai bang:")
        print('  python tai_androzoo.py --danh-sach "%s"' % ts.danh_sach)
    elif thieu_file:
        print()
        print("Them --mo-lai de tra trang thai va xuat %s." % TEN_DANH_SACH)

    # --- bo sung dong cho file mo coi ---
    if ts.bo_sung_dong and thieu_dong:
        try:
            from tai_androzoo import ghi_excel
        except ImportError as e:
            sys.exit("Khong import duoc ghi_excel tu tai_androzoo.py: %s" % e)
        tu_db = doc_db(con, thieu_dong)
        lan_chay = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        dong_moi = []
        thieu_db = 0
        for h in thieu_dong:
            d = tu_db.get(h, {})
            if not d:
                thieu_db += 1
            dong_moi.append({
                "Lan_chay": lan_chay,
                "SHA-256": h,
                "pkg_name": d.get("pkg_name") or "",
                "Kich_thuoc_byte": d.get("apk_size") or apk[h],
                "Moc_ngay": d.get("moc_ngay") or "",
                "dex_date": d.get("dex_date") or "",
                "vt_detection": d.get("vt_detection"),
                "vt_scan_date": d.get("vt_scan_date") or "",
                "markets": d.get("markets") or "",
                "Nhom": d.get("nhom") or "",
                "Duong_dan": os.path.join(ts.thu_muc, h + ".apk"),
                "Ghi_chu": "bo sung khi doi soat: file co tren dia, thieu dong",
            })
        kq = ghi_excel(os.path.dirname(os.path.abspath(ts.excel)), dong_moi, lan_chay)
        print()
        if kq:
            dd, them = kq
            print("Da ghi bo sung %d dong vao %s (ban cu nam trong _ban_cu)."
                  % (them, os.path.basename(dd)))
        if thieu_db:
            print("  %d dong khong tra duoc thong tin trong danh muc, chi co "
                  "SHA-256 va dung luong file." % thieu_db)
        # Danh dau lai trong DB cho khop thuc te.
        n = 0
        for i in range(0, len(thieu_dong), 500):
            lo = thieu_dong[i:i + 500]
            cur = con.execute(
                "UPDATE ungvien SET trang_thai='da_tai' WHERE sha256 IN (%s)"
                % ",".join("?" * len(lo)), lo)
            n += cur.rowcount
        con.commit()
        print("Da danh dau %d mau la 'da_tai' trong danh muc." % n)
    elif thieu_dong:
        print()
        print("Them --bo-sung-dong de ghi %d dong con thieu vao %s."
              % (len(thieu_dong), os.path.basename(ts.excel)))

    con.close()

    print()
    print("--- Sau khi lam xong het, ba con so nay phai bang nhau ---")
    print("  so file .apk trong thu muc apk")
    print("  so dong trong sheet %s cua %s" % (S_MAU, os.path.basename(ts.excel)))
    print("  so dong Nguon='AndroZoo' trong Manifest_thu_thap_APK.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
