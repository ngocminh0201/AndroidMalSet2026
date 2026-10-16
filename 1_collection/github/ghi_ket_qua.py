# -*- coding: utf-8 -*-
"""Module dung chung: doc va ghi so ket qua thu gom APK bang mot file Excel duy nhat.

Tu 30/8/2026 toan bo tuyen khong xuat CSV nua. Moi lan chay ghi NOI THEM vao
file <dich>\\Ket_qua_thu_gom.xlsx chu khong ghi de, nen khong con canh chay lai
lan hai lam mat so lieu lan mot nhu da xay ra hom 30/8.

Ba sheet trong file:
    Mau_thu_duoc      moi dong mot mau APK hop le, khoa la SHA-256, khong trung
    File_bi_loai      moi dong mot file bi loai, kem ly do
    Tong_hop_theo_kho bang gop theo tung kho GitHub, dung lai moi lan ghi

Cot Lan_chay ghi thoi diem cua lan thu gom sinh ra dong do, de truy nguoc.

Cac script khac dung module nay:
    thu_gom_tu_zip.py, thu_gom_github_apk.py   goi ghi_them(...)
    tong_hop_theo_kho.py                        goi doc(...) va gom_theo_kho(...)
"""

import collections
import datetime
import os
import re
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TEN_FILE = "Ket_qua_thu_gom.xlsx"
S_KQ, S_BQ, S_TH = "Mau_thu_duoc", "File_bi_loai", "Tong_hop_theo_kho"

COT_KQ = ["Lan_chay", "SHA-256", "Ten_file_goc", "Repo", "URL", "Dung_luong_byte",
          "Nam_classes_dex", "Nam_lon_nhat_trong_APK", "targetSdkVersion",
          "Duong_dan", "Ghi_chu"]
COT_BQ = ["Lan_chay", "Repo", "Ten_file", "Ly_do", "SHA256", "Trung_voi_kho"]
COT_SO = {"Dung_luong_byte", "Nam_classes_dex", "Nam_lon_nhat_trong_APK",
          "targetSdkVersion"}

FONT = "Times New Roman"
XANH = PatternFill("solid", fgColor="1F4E78")
CAM = PatternFill("solid", fgColor="FCE4D6")
XAM = PatternFill("solid", fgColor="D9E1F2")

# Moc nam do cong cu build ghi cung vao entry cua APK chu khong phai nam that:
# zipalign va mot so ban aapt dat 1980-01-01 hoac 1981-01-01, vai chuoi build
# khac dat 2008. Gap cac moc nay nghia la APK khong con dau vet thoi gian.
MOC_KHONG_THAT = {"", "0", "1980", "1981", "2008"}

# Bon kho cua dot 2 (thu muc GitHubMinh), dung de ghi "dot 1 / dot 2".
DOT2 = {"Pyran1/MalwareDatabase", "cryptwareapps/Malware-Database",
        "IHbib/goldoson-samples", "ytisf/theZoo"}


def duong_dan_file(dich):
    return os.path.join(dich, TEN_FILE)


def _doc_sheet(wb, ten, cot):
    if ten not in wb.sheetnames:
        return []
    ws = wb[ten]
    tieu_de = [c.value for c in ws[1]]
    ra = []
    for hang in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in hang):
            continue
        d = {t: ("" if v is None else v) for t, v in zip(tieu_de, hang)}
        ra.append({k: d.get(k, "") for k in cot})
    return ra


def doc(dich):
    """Tra ve (danh sach mau hop le, danh sach file bi loai). Chua co file thi tra ve rong."""
    dd = duong_dan_file(dich)
    if not os.path.exists(dd):
        return [], []
    wb = load_workbook(dd, read_only=True, data_only=True)
    kq = _doc_sheet(wb, S_KQ, COT_KQ)
    bq = _doc_sheet(wb, S_BQ, COT_BQ)
    wb.close()
    return kq, bq


def gom_theo_kho(kq, bq):
    """Gop theo tung kho, dem theo SHA-256 chu khong theo so dong.

    Mot hash da nam trong sheet mau hop le thi khong bao gio bi tinh la khong
    hop le nua: lan chay sau se gap lai chinh nhung mau da lay lan truoc va ghi
    chung vao phan bi loai voi ly do "trung voi mau da co". Do la mot luot gap
    lai cua cung mot file, khong phai mot file hong.
    """
    kho = collections.OrderedDict()

    def lay(ten):
        if ten not in kho:
            kho[ten] = {"url": "", "hop_le": 0, "trung_dot": 0, "trung_da_co": 0,
                        "moc_khong_that": 0, "moc_cu": 0, "khong_giai_nen": 0,
                        "khac": 0, "lap_lai": 0, "nam_cu": collections.Counter()}
        return kho[ten]

    da_hop_le = set()
    for x in kq:
        h = str(x.get("SHA-256", "")).strip().lower()
        if not h or h in da_hop_le:
            continue
        da_hop_le.add(h)
        a = lay(x["Repo"])
        a["hop_le"] += 1
        a["url"] = x.get("URL") or a["url"]

    da_loai = set()
    for x in bq:
        h = str(x.get("SHA256", "")).strip().lower()
        ly_do = str(x.get("Ly_do", ""))
        a = lay(x["Repo"])
        if not h:
            if "khong giai nen" in ly_do.lower():
                a["khong_giai_nen"] += 1
            continue
        if h in da_hop_le:
            a["lap_lai"] += 1
            continue
        if h in da_loai:
            continue
        da_loai.add(h)
        m = re.search(r"moc thoi gian (\d{4})", ly_do)
        if m:
            if m.group(1) in MOC_KHONG_THAT:
                a["moc_khong_that"] += 1
            else:
                a["moc_cu"] += 1
                a["nam_cu"][m.group(1)] += 1
        elif "trung ban da lay trong dot nay" in ly_do:
            a["trung_dot"] += 1
        elif "trung voi mau da co" in ly_do:
            a["trung_da_co"] += 1
        else:
            a["khac"] += 1
    return kho


def viet_ghi_chu(ten_kho, a):
    phan = []
    dot = "đợt 2" if ten_kho in DOT2 else "đợt 1"
    if a["trung_dot"]:
        phan.append("%d file trùng hash với mẫu đã lấy trong cùng đợt" % a["trung_dot"])
    if a["trung_da_co"]:
        phan.append("%d file trùng với mẫu đã có trong manifest hoặc kho apk" % a["trung_da_co"])
    if a["moc_cu"]:
        nam = ", ".join("%s: %d" % (n, s) for n, s in sorted(a["nam_cu"].items()))
        phan.append("%d file có mốc thời gian trong APK trước ngưỡng năm đã đặt (%s)"
                    % (a["moc_cu"], nam))
    if a["moc_khong_that"]:
        phan.append("%d file bị loại vì mốc 1980/1981/2008 — mốc do công cụ build ghi cứng "
                    "chứ không phải năm thật của mẫu, cần chạy lại để lấy về"
                    % a["moc_khong_that"])
    if a["khong_giai_nen"]:
        phan.append("%d file nén không giải nén được, chưa rõ bên trong có APK hay không, "
                    "không tính vào cột tổng" % a["khong_giai_nen"])
    if a["khac"]:
        phan.append("%d file bị loại vì lý do khác" % a["khac"])
    if a["lap_lai"]:
        phan.append("%d lượt gặp lại mẫu đã thu (trùng hash giữa các kho hoặc do chạy lại) "
                    "— cùng một file nên không tính vào cột không hợp lệ" % a["lap_lai"])
    if not phan:
        return "Thu gom %s. Không có file nào bị loại." % dot
    return "Thu gom %s. " % dot + "; ".join(phan) + "."


def _ke_bang(ws, cot, du_lieu, to_cam=None):
    ws.append(cot)
    for o in ws[1]:
        o.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        o.fill = XANH
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for d in du_lieu:
        ws.append([int(d[k]) if k in COT_SO and str(d[k]).lstrip("-").isdigit() else d[k]
                   for k in cot])

    for hang in ws.iter_rows(min_row=2):
        for o in hang:
            o.font = Font(name=FONT, size=10)
            o.alignment = Alignment(vertical="top")
        if to_cam and to_cam(hang):
            for o in hang:
                o.fill = CAM

    for i, ten in enumerate(cot, start=1):
        rong = max([len(str(ten))] + [len(str(d[ten])) for d in du_lieu[:500]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(rong, 9), 58)
    if du_lieu:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cot)), len(du_lieu) + 1)
    ws.freeze_panes = "B2"


COT_TH = [("STT", 6), ("Link repo", 52), ("Tổng số file APK trong kho", 16),
          ("Số file hợp lệ", 12), ("Số file không hợp lệ", 14), ("Ghi chú", 88)]


def ke_bang_tong_hop(ws, kho, chu_thich):
    """Dung bang tong hop theo kho. Dung chung cho file Ket_qua_thu_gom.xlsx va
    cho sheet Tong_hop_theo_kho trong manifest."""
    ws["A1"] = "TỔNG HỢP MẪU APK THU THẬP THEO TỪNG KHO GITHUB"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws.merge_cells("A1:F1")

    hang = 3
    for t in chu_thich:
        o = ws.cell(row=hang, column=1, value=t)
        o.font = Font(name=FONT, size=9, italic=True,
                      color="C00000" if t.startswith("CẦN XỬ LÝ") else "404040")
        o.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=6)
        ws.row_dimensions[hang].height = 26 if len(t) < 200 else 40
        hang += 1

    hang += 1
    hang_tieu_de = hang
    for i, (ten, rong) in enumerate(COT_TH, start=1):
        o = ws.cell(row=hang, column=i, value=ten)
        o.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        o.fill = XANH
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = rong
    ws.row_dimensions[hang].height = 32

    thu_tu = sorted(kho.items(), key=lambda kv: (kv[0] in DOT2, -kv[1]["hop_le"], kv[0].lower()))
    hang_dau = hang + 1
    for stt, (ten_kho, a) in enumerate(thu_tu, start=1):
        hang += 1
        khong_hop_le = (a["trung_dot"] + a["trung_da_co"] + a["moc_khong_that"]
                        + a["moc_cu"] + a["khac"])
        ws.cell(row=hang, column=1, value=stt)
        ws.cell(row=hang, column=2, value=a["url"] or "https://github.com/" + ten_kho)
        ws.cell(row=hang, column=3, value="=D%d+E%d" % (hang, hang))
        ws.cell(row=hang, column=4, value=a["hop_le"])
        ws.cell(row=hang, column=5, value=khong_hop_le)
        ws.cell(row=hang, column=6, value=viet_ghi_chu(ten_kho, a))
        for i in range(1, 7):
            o = ws.cell(row=hang, column=i)
            o.font = Font(name=FONT, size=10)
            o.alignment = Alignment(horizontal="center" if i in (1, 3, 4, 5) else "left",
                                    vertical="top", wrap_text=(i == 6))
            if a["hop_le"] == 0:
                o.fill = CAM
        ws.row_dimensions[hang].height = 30
    hang_cuoi = hang

    hang += 1
    ws.cell(row=hang, column=2, value="TỔNG CỘNG %d kho" % len(thu_tu))
    for i, chu in ((3, "C"), (4, "D"), (5, "E")):
        ws.cell(row=hang, column=i, value="=SUM(%s%d:%s%d)" % (chu, hang_dau, chu, hang_cuoi))
    ws.cell(row=hang, column=6,
            value="Số mẫu hợp lệ là số APK duy nhất theo SHA-256, đã khử trùng lặp giữa các "
                  "kho và với manifest.")
    for i in range(1, 7):
        o = ws.cell(row=hang, column=i)
        o.font = Font(name=FONT, size=11, bold=True)
        o.fill = XAM
        o.alignment = Alignment(horizontal="center" if i in (1, 3, 4, 5) else "left",
                                vertical="center", wrap_text=(i == 6))
    ws.row_dimensions[hang].height = 30
    ws.freeze_panes = ws.cell(row=hang_dau, column=1).coordinate
    ws.auto_filter.ref = "A%d:F%d" % (hang_tieu_de, hang_cuoi)


def chu_thich_chuan(dich, kho, so_file_apk):
    """dich o day chi dung de hien thi trong phan chu thich."""
    tong_hop_le = sum(a["hop_le"] for a in kho.values())
    tong_moc = sum(a["moc_khong_that"] for a in kho.values())
    duong_dan_hien = os.path.join(dich, TEN_FILE).replace("/", "\\")
    ra = ["Số liệu lấy từ %s, lập ngày %s. Đối chiếu: thư mục %s\\apk có %d file .apk, "
          "tổng số mẫu hợp lệ đếm được là %d — %s."
          % (duong_dan_hien, datetime.date.today().strftime("%d/%m/%Y"),
             dich, so_file_apk, tong_hop_le,
             "khớp" if so_file_apk == tong_hop_le
             else "LỆCH %d file" % abs(so_file_apk - tong_hop_le)),
          "Cột \"Tổng số file APK trong kho\" là số file đã được nhận diện là APK theo nội "
          "dung (ZIP hợp lệ, bên trong có AndroidManifest.xml và classes*.dex), tính cả file "
          "nằm trong các file nén lồng. Bằng số hợp lệ cộng số không hợp lệ.",
          "Đợt 1 (12 kho đầu) chạy không đặt ngưỡng năm nên không có file nào bị loại vì mốc "
          "thời gian. Đợt 2 (4 kho trong thư mục GitHubMinh) chạy với --nam-toi-thieu 2020."]
    if so_file_apk and so_file_apk != tong_hop_le:
        ra.append("CẦN XỬ LÝ: số file trong thư mục apk không bằng số dòng trong sheet %s. "
                  "Không xoá file thừa, hãy tìm phần ghi nhận bị thiếu rồi bổ sung." % S_KQ)
    if tong_moc:
        ra.append("CẦN XỬ LÝ: %d file bị loại vì mốc 1980, 1981 hoặc 2008. Đây là mốc do công "
                  "cụ build ghi cứng chứ không phải năm thật của mẫu, nên đó là loại nhầm. "
                  "Chạy lại phần thu gom để lấy về." % tong_moc)
    return ra


def dem_file_apk(dich):
    thu_muc = os.path.join(dich, "apk")
    if not os.path.isdir(thu_muc):
        return 0
    return len([t for t in os.listdir(thu_muc) if re.fullmatch(r"[0-9a-fA-F]{64}\.apk", t)])


def ghi_them(dich, ket_qua_moi, bo_qua_moi, nhan_lan_chay="", nhan_dich=""):
    """Ghi noi them ket qua cua mot lan chay vao Ket_qua_thu_gom.xlsx.

    ket_qua_moi / bo_qua_moi la danh sach dict theo dung ten cot COT_KQ / COT_BQ
    (khong can cot Lan_chay, ham tu dien). Dong nao co SHA-256 da nam trong sheet
    thi khong them lai. Tra ve (so dong hop le them moi, tong so mau hop le).
    """
    dd = duong_dan_file(dich)
    nhan = nhan_lan_chay or datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    kq, bq = doc(dich)

    if os.path.exists(dd):
        ban_sao = os.path.join(dich, "_ban_cu")
        os.makedirs(ban_sao, exist_ok=True)
        shutil.copy2(dd, os.path.join(
            ban_sao, "Ket_qua_thu_gom_%s.xlsx"
            % datetime.datetime.now().strftime("%d-%m-%Y_%Hh%Mm%Ss")))

    da_co = {str(x["SHA-256"]).strip().lower() for x in kq}
    them = 0
    for x in ket_qua_moi:
        h = str(x.get("SHA-256", "")).strip().lower()
        if not h or h in da_co:
            continue
        da_co.add(h)
        d = {k: x.get(k, "") for k in COT_KQ}
        d["Lan_chay"] = nhan
        kq.append(d)
        them += 1
    for x in bo_qua_moi:
        d = {k: x.get(k, "") for k in COT_BQ}
        d["Lan_chay"] = nhan
        bq.append(d)

    kho = gom_theo_kho(kq, bq)
    wb = Workbook()
    wb.remove(wb.active)
    _ke_bang(wb.create_sheet(S_KQ), COT_KQ, kq)
    _ke_bang(wb.create_sheet(S_BQ), COT_BQ, bq,
             to_cam=lambda hang: "khong giai nen" in str(hang[3].value).lower())
    ke_bang_tong_hop(wb.create_sheet(S_TH), kho,
                     chu_thich_chuan(nhan_dich or dich, kho, dem_file_apk(dich)))
    wb.save(dd)
    return them, len(kq)
