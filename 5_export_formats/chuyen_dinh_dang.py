# -*- coding: utf-8 -*-
"""chuyen_dinh_dang.py - chuyen ma tran dac trung apk_ml_dataset.csv sang cac
dinh dang dung truc tiep cho cac cong cu hoc may pho bien.

    arff    ARFF dang thua (sparse), mo bang Weka (Explorer, Experimenter, CLI)
    libsvm  LibSVM / SVMlight, doc bang scikit-learn (load_svmlight_file),
            XGBoost, LightGBM, LIBSVM, Weka (LibSVMLoader)
    npz     ma tran thua SciPy (CSR, float32) + nhan NumPy, nap vao scikit-learn,
            Keras/TensorFlow, PyTorch trong Anaconda, Jupyter Notebook, Google Colab

File CSV goc co 20.035 cot (file_name, source, sha256 va 20.032 dac trung) va
hon 120 nghin dong, nang khoang 4,9 GB. Script doc theo khoi bang pandas nen
khong can nap ca file vao RAM; ARFF va LibSVM duoc ghi dan tung khoi, chi NPZ
can giu ma tran thua trong RAM (vai GB voi bo du lieu day du).

Nhan lay tu mot file rieng (CSV hoac XLSX) co cot SHA-256 va cot nhan, ghep
theo sha256. Mac dinh bo cac mau khong co nhan (vi du mau chua co bao cao
VirusTotal); --giu-khong-nhan de giu lai voi nhan "?" (ARFF) hoac -1 (NPZ).

Vi du:
    python chuyen_dinh_dang.py --csv apk_ml_dataset.csv --output xuat ^
        --nhan Ket_qua_GanNhan.xlsx --cot-nhan Nhan_bai_bao --dinh-dang arff libsvm npz

    # bai toan nhi phan Benign/Malware
    python chuyen_dinh_dang.py --csv apk_ml_dataset.csv --output xuat_nhi_phan ^
        --nhan Ket_qua_GanNhan.xlsx --cot-nhan Nhan_bai_bao --nhi-phan --dinh-dang arff npz

    # chay thu tren 2.000 dong dau
    python chuyen_dinh_dang.py --csv apk_ml_dataset.csv --output thu --gioi-han 2000
"""

import argparse
import csv
import os
import sys
import time

import numpy as np
import pandas as pd
from scipy import sparse

COT_ID = ["file_name", "source", "sha256"]
BENIGN = "Benign"
CAC_DINH_DANG = ("arff", "libsvm", "npz")


# ------------------------------------------------------------------ doc nhan

def tim_cot(tieu_de, ung_vien, ten_hien_thi):
    thap = {str(c).strip().lower(): c for c in tieu_de if c is not None}
    for u in ung_vien:
        if u and u.lower() in thap:
            return thap[u.lower()]
    sys.exit("Khong tim thay cot %s trong file nhan. Cac cot dang co: %s"
             % (ten_hien_thi, ", ".join(str(c) for c in tieu_de)))


def doc_nhan(duong_dan, cot_sha, cot_nhan):
    """Tra ve dict sha256 (chu thuong) -> nhan."""
    if duong_dan.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(duong_dan, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        dong = ws.iter_rows(values_only=True)
        tieu_de = list(next(dong))
        c_sha = tim_cot(tieu_de, [cot_sha, "sha256", "SHA-256", "SHA256"], "SHA-256")
        c_nhan = tim_cot(tieu_de, [cot_nhan, "label", "Nhan_bai_bao", "Nhan"], "nhan")
        i_sha, i_nhan = tieu_de.index(c_sha), tieu_de.index(c_nhan)
        nhan = {}
        for r in dong:
            s, n = r[i_sha], r[i_nhan]
            if s and n not in (None, ""):
                nhan[str(s).strip().lower()] = str(n).strip()
        wb.close()
    else:
        df = pd.read_csv(duong_dan, dtype=str, keep_default_na=False)
        c_sha = tim_cot(list(df.columns), [cot_sha, "sha256", "SHA-256", "SHA256"], "SHA-256")
        c_nhan = tim_cot(list(df.columns), [cot_nhan, "label", "Nhan_bai_bao", "Nhan"], "nhan")
        df = df[(df[c_sha] != "") & (df[c_nhan] != "")]
        nhan = dict(zip(df[c_sha].str.strip().str.lower(), df[c_nhan].str.strip()))
    print("Doc nhan: %d mau tu %s (cot '%s')" % (len(nhan), duong_dan, c_nhan))
    return nhan


# ------------------------------------------------------------------ dinh dang so

def dinh_dang_so(gia_tri):
    """float32 -> chuoi gon: so nguyen ghi khong co phan thap phan."""
    return np.char.mod("%.7g", gia_tri.astype(np.float64))


def nhay_arff(ten):
    """Bao ten thuoc tinh/gia tri trong dau nhay don theo quy uoc ARFF."""
    return "'" + str(ten).replace("\\", "\\\\").replace("'", "\\'") + "'"


# ------------------------------------------------------------------ cac bo ghi

class GhiARFF:
    def __init__(self, duong_dan, ten_cot, lop):
        self.f = open(duong_dan, "w", encoding="utf-8", newline="\n")
        self.so_cot = len(ten_cot)
        self.lop = lop
        self.f.write("% AndroidMalSet2026 - ma tran dac trung dang ARFF thua (sparse)\n")
        self.f.write("% Thu tu dong trung voi sample_ids.csv\n")
        self.f.write("@RELATION AndroidMalSet2026\n\n")
        for c in ten_cot:
            self.f.write("@ATTRIBUTE %s NUMERIC\n" % nhay_arff(c))
        if lop is not None:
            self.f.write("@ATTRIBUTE class {%s}\n" % ",".join(nhay_arff(x) for x in lop))
        self.f.write("\n@DATA\n")

    def ghi(self, X, nhan):
        for i in range(X.shape[0]):
            a, b = X.indptr[i], X.indptr[i + 1]
            phan = [str(j) + " " + v for j, v in zip(X.indices[a:b].tolist(),
                                                     dinh_dang_so(X.data[a:b]).tolist())]
            if self.lop is not None:
                n = nhan[i]
                phan.append("%d %s" % (self.so_cot, nhay_arff(n) if n is not None else "?"))
            self.f.write("{" + ",".join(phan) + "}\n")

    def dong(self):
        self.f.close()


class GhiLibSVM:
    def __init__(self, duong_dan, lop):
        self.f = open(duong_dan, "w", encoding="utf-8", newline="\n")
        self.ma = {x: k for k, x in enumerate(lop)} if lop is not None else None

    def ghi(self, X, nhan):
        for i in range(X.shape[0]):
            a, b = X.indptr[i], X.indptr[i + 1]
            y = -1 if self.ma is None or nhan[i] is None else self.ma[nhan[i]]
            phan = [str(j + 1) + ":" + v for j, v in zip(X.indices[a:b].tolist(),
                                                         dinh_dang_so(X.data[a:b]).tolist())]
            self.f.write(str(y) + (" " + " ".join(phan) if phan else "") + "\n")

    def dong(self):
        self.f.close()


# ------------------------------------------------------------------ chuong trinh chinh

def main():
    p = argparse.ArgumentParser(description="Chuyen apk_ml_dataset.csv sang ARFF, LibSVM, NPZ.")
    p.add_argument("--csv", required=True, help="file apk_ml_dataset.csv")
    p.add_argument("--output", required=True, help="thu muc ghi ket qua")
    p.add_argument("--dinh-dang", nargs="+", default=["arff", "libsvm", "npz"],
                   choices=CAC_DINH_DANG, help="mot hoac nhieu dinh dang can xuat")
    p.add_argument("--nhan", default="", help="file nhan (CSV hoac XLSX), ghep theo sha256")
    p.add_argument("--cot-sha", default="sha256", help="ten cot SHA-256 trong file nhan")
    p.add_argument("--cot-nhan", default="label", help="ten cot nhan trong file nhan")
    p.add_argument("--nhi-phan", action="store_true",
                   help="doi nhan ve hai lop Benign/Malware (moi nhan khac Benign la Malware)")
    p.add_argument("--chi-ma-doc", action="store_true",
                   help="chi giu mau ma doc (bai toan phan loai bon nhom)")
    p.add_argument("--giu-khong-nhan", action="store_true", help="giu ca mau khong co nhan")
    p.add_argument("--khoi", type=int, default=2000, help="so dong doc moi khoi (mac dinh 2000)")
    p.add_argument("--gioi-han", type=int, default=0, help="chi doc N dong dau, de chay thu")
    a = p.parse_args()

    os.makedirs(a.output, exist_ok=True)
    bat_dau = time.time()

    tieu_de = pd.read_csv(a.csv, nrows=0).columns.tolist()
    thieu = [c for c in COT_ID if c not in tieu_de]
    if thieu:
        sys.exit("File CSV thieu cot: %s" % ", ".join(thieu))
    ten_cot = [c for c in tieu_de if c not in COT_ID]
    print("CSV: %s - %d dac trung" % (a.csv, len(ten_cot)))

    bang_nhan = doc_nhan(a.nhan, a.cot_sha, a.cot_nhan) if a.nhan else None
    if bang_nhan is not None and a.nhi_phan:
        bang_nhan = {k: (BENIGN if v == BENIGN else "Malware") for k, v in bang_nhan.items()}
    if bang_nhan is not None and a.chi_ma_doc:
        bang_nhan = {k: v for k, v in bang_nhan.items() if v != BENIGN}
    lop = sorted(set(bang_nhan.values())) if bang_nhan is not None else None
    if lop is not None:
        print("Cac lop: %s" % ", ".join(lop))
    elif "arff" in a.dinh_dang or "libsvm" in a.dinh_dang:
        print("Canh bao: khong co --nhan, ARFF se khong co thuoc tinh lop, LibSVM ghi nhan -1.")

    with open(os.path.join(a.output, "feature_names.txt"), "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(ten_cot) + "\n")
    f_id = open(os.path.join(a.output, "sample_ids.csv"), "w", encoding="utf-8", newline="")
    w_id = csv.writer(f_id)
    w_id.writerow(COT_ID + ["label"])

    bo_ghi = []
    if "arff" in a.dinh_dang:
        bo_ghi.append(GhiARFF(os.path.join(a.output, "AndroidMalSet2026.arff"), ten_cot, lop))
    if "libsvm" in a.dinh_dang:
        bo_ghi.append(GhiLibSVM(os.path.join(a.output, "AndroidMalSet2026.libsvm"), lop))
    can_giu = "npz" in a.dinh_dang
    khoi_X, tat_ca_nhan = [], []

    kieu = {c: "float32" for c in ten_cot}
    kieu.update({c: str for c in COT_ID})
    doc = pd.read_csv(a.csv, dtype=kieu, chunksize=a.khoi, keep_default_na=False,
                      na_values=[""], nrows=a.gioi_han or None)
    da_doc = giu = bo = 0
    for df in doc:
        da_doc += len(df)
        if bang_nhan is not None:
            nhan = [bang_nhan.get(str(s).strip().lower()) for s in df["sha256"]]
            mask = np.array([n is not None for n in nhan]) if not a.giu_khong_nhan \
                else np.ones(len(df), dtype=bool)
        else:
            nhan = [None] * len(df)
            mask = np.ones(len(df), dtype=bool)
        bo += int((~mask).sum())
        if not mask.any():
            continue
        df = df.loc[mask]
        nhan = [n for n, m in zip(nhan, mask) if m]
        A = df[ten_cot].to_numpy(dtype=np.float32, na_value=0.0)
        X = sparse.csr_matrix(A)
        del A
        for g in bo_ghi:
            g.ghi(X, nhan)
        for (_, r), n in zip(df[COT_ID].iterrows(), nhan):
            w_id.writerow([r["file_name"], r["source"], r["sha256"], "" if n is None else n])
        if can_giu:
            khoi_X.append(X)
            tat_ca_nhan.extend(nhan)
        giu += X.shape[0]
        print("  da doc %d dong, giu %d, bo %d (%.0f giay)" % (da_doc, giu, bo, time.time() - bat_dau))

    for g in bo_ghi:
        g.dong()
    f_id.close()

    if can_giu and khoi_X:
        X = sparse.vstack(khoi_X, format="csr", dtype=np.float32)
        if "npz" in a.dinh_dang:
            sparse.save_npz(os.path.join(a.output, "X.npz"), X, compressed=True)
            if lop is not None:
                ma = {x: k for k, x in enumerate(lop)}
                y = np.array([ma.get(n, -1) if n is not None else -1 for n in tat_ca_nhan], dtype=np.int16)
                np.save(os.path.join(a.output, "y.npy"), y)
                with open(os.path.join(a.output, "classes.txt"), "w", encoding="utf-8", newline="\n") as f:
                    f.write("\n".join(lop) + "\n")

    print("Xong: %d mau, %d dac trung, %d mau bi bo vi khong co nhan. Ket qua o %s (%.0f giay)"
          % (giu, len(ten_cot), bo, os.path.abspath(a.output), time.time() - bat_dau))


if __name__ == "__main__":
    main()
