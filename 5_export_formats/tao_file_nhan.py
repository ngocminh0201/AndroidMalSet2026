"""
Build the public label file (labels.csv) of AndroidMalSet2026 and the checksum list
for the Zenodo deposit.

Only fields that may be redistributed are written: SHA-256, source, labels, the number
of flagging VirusTotal engines, the total number of engines, the label confidence,
the labelling method/model and the first VirusTotal submission date.
Per-engine results, VirusTotal family/threat names and Gemini free-text reasons are
NOT written (VirusTotal terms of service; reasons are in Vietnamese).

Usage (Windows, from any folder):
  python tao_file_nhan.py --gan-nhan <folder>\\Ket_qua_GanNhan.xlsx ^
                          --virustotal <folder>\\Ket_qua_VirusTotal.xlsx ^
                          --ma-tran <folder>\\apk_ml_dataset.csv ^
                          --out <release_folder>
Optional: --them-file <file> (repeatable) adds more files to SHA256SUMS.txt
          (e.g. feature_vocabulary.csv, dynamic_features.csv).
"""
import argparse, csv, hashlib, os, sys, shutil
from collections import Counter
from datetime import datetime
import openpyxl

CATS = ["Riskware", "Adware", "SMS malware", "Banking malware"]

def read_labels(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Nhan_cuoi"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it)); ix = {h: i for i, h in enumerate(head)}
    out = {}
    for r in it:
        sha = (r[ix["SHA-256"]] or "").strip().lower()
        if not sha: continue
        out[sha] = dict(
            label=r[ix["Nhan_bai_bao"]], pos=r[ix["So_engine_phat_hien"]], tot=r[ix["Tong_so_engine"]],
            conf=r[ix["Do_tin_cay"]], method=r[ix["Nguon_nhan"]], model=r[ix["Model"]])
    return out

def read_first_seen(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Ket_qua_VT"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it)); ix = {h: i for i, h in enumerate(head)}
    out = {}
    for r in it:
        sha = (r[ix["SHA-256"]] or "").strip().lower()
        d = r[ix["Lan_gui_dau_tien"]]
        if sha and d:
            try: out[sha] = datetime.strptime(str(d).strip(), "%d/%m/%Y").strftime("%Y-%m-%d")
            except ValueError: out[sha] = ""
    return out

def read_matrix_keys(path):
    """Stream only the first three columns (file_name, source, sha256) of the 4.89 GB CSV."""
    keys = {}
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        head = [h.strip().lstrip("\ufeff").strip('"') for h in f.readline().lstrip("\ufeff").rstrip("\r\n").split(",")[:3]]
        assert head == ["file_name", "source", "sha256"], head
        for n, line in enumerate(f, 1):
            p = line.split(",", 3)
            sha = (p[2] or p[0].replace(".apk", "")).strip().lower()
            keys[sha] = p[1].strip()
            if n % 20000 == 0: print(f"  read {n:,} rows of the feature matrix", flush=True)
    return keys

def sha256_of(path, buf=1 << 24):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(buf)
            if not b: break
            h.update(b)
    return h.hexdigest()

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--gan-nhan", required=True)
    ap.add_argument("--virustotal", required=True)
    ap.add_argument("--ma-tran", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--them-file", action="append", default=[])
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    print("tao_file_nhan.py v3 (BOM-safe)")
    print("Reading labels ..."); lab = read_labels(a.gan_nhan)
    print("Reading VirusTotal first-submission dates ..."); first = read_first_seen(a.virustotal)
    print("Reading the keys of the feature matrix ..."); mat = read_matrix_keys(a.ma_tran)
    rows = []
    for sha, src in mat.items():
        if sha not in lab: continue            # the 214 samples without a VirusTotal report
        L = lab[sha]
        if L["label"] not in ["Benign"] + CATS: continue
        rows.append(dict(
            sha256=sha, source=src,
            binary_label="Benign" if L["label"] == "Benign" else "Malware",
            category_label=L["label"],
            vt_positives=int(L["pos"]), vt_total=int(L["tot"]),
            label_confidence=L["conf"],
            label_method="zero_engine_rule" if L["label"] == "Benign" else "gemini",
            gemini_model=L["model"] or "",
            vt_first_submission_date=first.get(sha, "")))
    rows.sort(key=lambda r: r["sha256"])
    out_csv = os.path.join(a.out, "labels.csv")
    with open(out_csv, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
    c = Counter(r["category_label"] for r in rows)
    print(f"labels.csv: {len(rows):,} rows x {len(rows[0])} columns, {os.path.getsize(out_csv)/1e6:.2f} MB")
    print("  per class:", dict(c))
    expected = {"Benign": 62919, "Riskware": 34969, "Adware": 14005, "SMS malware": 6070, "Banking malware": 2160}
    if len(rows) != 120123 or dict(c) != expected:
        print("!! WARNING: counts differ from the manuscript (120,123 samples; Table 3). Check before uploading.")
    else:
        print("  OK: counts match the manuscript (120,123 samples, Table 3).")
    files = [out_csv, a.ma_tran] + a.them_file
    with open(os.path.join(a.out, "SHA256SUMS.txt"), "w", encoding="utf-8") as f:
        for p in files:
            print(f"Hashing {p} ..."); f.write(f"{sha256_of(p)}  {os.path.basename(p)}\n")
            print(f"  size: {os.path.getsize(p):,} bytes ({os.path.getsize(p)/1e9:.2f} GB)")
    print("Done. Upload labels.csv, SHA256SUMS.txt, README.md and the other data files to Zenodo.")

if __name__ == "__main__":
    main()
